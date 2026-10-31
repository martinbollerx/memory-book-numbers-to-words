"""Turn the Brysbaert concreteness ratings into build/concreteness.json.

Source: Brysbaert, Warriner & Kuperman (2014), "Concreteness ratings for 40
thousand generally known English word lemmas", Behavior Research Methods.
Supplementary file 13428_2013_403_MOESM1_ESM.xlsx from
https://link.springer.com/article/10.3758/s13428-013-0403-5

Usage: python build/export_ratings.py path/to/13428_2013_403_MOESM1_ESM.xlsx

Writes {"word": rating} where rating runs from 1 (abstract) to 5 (concrete),
single words only. The spreadsheet and this json are not committed; only the
ranking they produce (a combined score per word) ends up in data/index.json.

Requires: pip install openpyxl
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

source = Path(sys.argv[1])
ws = load_workbook(source, read_only=True)["Sheet1"]

ratings = {}
for word, bigram, conc, *_ in list(ws.iter_rows(values_only=True))[1:]:
    if bigram == 0 and isinstance(word, str) and isinstance(conc, (int, float)):
        ratings[word.lower()] = round(float(conc), 2)

out = Path(__file__).parent / "concreteness.json"
out.write_text(json.dumps(ratings), encoding="utf-8")
print(f"{len(ratings)} words written to {out}")
